"""Slide themes (layouts).

A Theme owns everything visual: colors, fonts, title-slide composition and
the recurring chrome of content slides (background, title, footer, logo,
references). The Deck in factory.py only places content inside the
rectangle a theme hands back.

This package ships one neutral theme, `plain`, and no logos. Branded themes
live in separate packages: subclass Theme, decorate it with @register_theme,
and declare the module under the `slide_factory.themes` entry-point group in
that package's pyproject.toml — Deck(theme="my_lab") then finds it by name:

    [project.entry-points."slide_factory.themes"]
    my_lab = "my_themes"

    from slide_factory.themes import Theme, register_theme

    @register_theme
    class MyLab(Theme):
        name = "my_lab"
        primary = RGBColor(0x12, 0x34, 0x56)
        default_logo = Path(__file__).parent / "assets" / "my_lab.png"

Logos: a theme may point `default_logo` at an image; it is used on every
slide. On dark brand surfaces the artwork is converted to a white,
transparent-background version automatically, so a plain downloaded logo
(even a JPEG on white) works.
"""
from __future__ import annotations

import io
from pathlib import Path

from pptx.dml.color import RGBColor

from .draw import (BLACK, IN, SLIDE_H, SLIDE_W, WHITE, Rect, arrow, box,
                   circle, fit_picture, para, run, set_fill_alpha,
                   set_picture_alpha, textbox)

W, H = SLIDE_W, SLIDE_H

THEMES: dict[str, type["Theme"]] = {}


def register_theme(cls: type["Theme"]) -> type["Theme"]:
    THEMES[cls.name] = cls
    return cls


_plugins_loaded = False


def load_theme_plugins() -> None:
    """Import every installed module declared under the `slide_factory.themes`
    entry-point group; importing it runs its @register_theme decorators."""
    global _plugins_loaded
    if _plugins_loaded:
        return
    _plugins_loaded = True
    from importlib.metadata import entry_points
    for ep in entry_points(group="slide_factory.themes"):
        try:
            ep.load()
        except Exception as exc:  # a broken theme package must not break the engine
            import warnings
            warnings.warn(f"slide_factory theme plugin {ep.name!r} failed to load: {exc}")


def make_theme(theme: "str | Theme | type[Theme]") -> "Theme":
    if isinstance(theme, Theme):
        return theme
    if isinstance(theme, type) and issubclass(theme, Theme):
        return theme()
    if str(theme) not in THEMES:
        load_theme_plugins()
    try:
        return THEMES[str(theme)]()
    except KeyError:
        raise ValueError(
            f"Unknown theme {theme!r}. Available: {sorted(THEMES)}. Branded themes "
            "come from separately installed theme packages.") from None


def mix(a: RGBColor, b: RGBColor, frac: float) -> RGBColor:
    """Blend a toward b; frac=0 -> a, frac=1 -> b. Used to 'dim' on solid bg."""
    return RGBColor(*(int(round(ca + (cb - ca) * frac)) for ca, cb in zip(a, b)))


def logo_white(path) -> io.BytesIO:
    """White-on-transparent version of a logo, for dark brand surfaces.

    Near-white background becomes transparent, all remaining artwork becomes
    flat white. Artwork that is already light-on-transparent is kept as is.
    """
    from PIL import Image, ImageChops, ImageStat

    im = Image.open(path).convert("RGBA")
    r, g, b, a = im.split()
    if im.getextrema()[3][0] < 250:                      # real transparency
        mask = a.point(lambda v: 255 if v > 128 else 0)
        if ImageStat.Stat(im.convert("L"), mask).mean[0] > 200:
            out = im                                      # already white art
        else:
            out = None
    else:
        out = None
    if out is None:
        minc = ImageChops.darker(ImageChops.darker(r, g), b)
        art = minc.point(lambda v: 0 if v > 235 else 255)
        out = Image.new("RGBA", im.size, (255, 255, 255, 0))
        out.putalpha(ImageChops.darker(art, a))
    buf = io.BytesIO()
    out.save(buf, "PNG")
    buf.seek(0)
    return buf


_mix, _logo_white = mix, logo_white   # pre-split private names, still imported by decks


class Theme:
    """Base theme: tokens + default composition. Subclasses mostly override tokens."""

    name = "base"

    # -- colors ----------------------------------------------------------
    primary = RGBColor(0x33, 0x33, 0x33)      # brand color (titles, footer band)
    on_primary = WHITE                        # text/logo color on primary
    bg = WHITE
    text = RGBColor(0x1A, 0x1A, 0x1A)
    muted = RGBColor(0x5A, 0x5A, 0x5A)
    title_authors_color = WHITE               # author line on the title slide
    title_subtitle_color = None               # None -> derived from on_primary
    takeaway_color = RGBColor(0x1A, 0x1A, 0x1A)
    accent = None                             # secondary brand color for the
    #                                           conceptual layouts (chain nodes,
    #                                           headlines, card borders);
    #                                           None -> takeaway_color

    def accent_color(self) -> RGBColor:
        return self.accent or self.takeaway_color

    # -- fonts / sizes (pt) ----------------------------------------------
    title_font = "Arial"
    body_font = "Arial"
    t_title = 38          # title slide
    c_title = 28          # content slide title
    body = 15
    sub = 13
    caption = 11
    takeaway = 17
    refs = 8
    refs_single = 10       # a lone citation: unnumbered, snug pill, this size
    bullet_char = "•"
    sub_bullet_char = "–"  # level >= 1 bullets

    # -- geometry / chrome -----------------------------------------------
    margin = IN(0.55)
    footer_h = IN(0.62)
    c_title_x = 0.42       # content-slide title position (inches)
    c_title_y = 0.35
    refs_boxed = True            # white reference pill (house style);
    #                              False = plain text on the footer band
    org_lines = ("ORGANISATION",)  # placeholder logotype when no logo image
    default_logo = None            # theme-level logo used when none is passed
    logo_has_wordmark = False      # logo artwork already contains the org name
    #                                -> org_mark/org_mark_centered draw no text

    # ------------------------------------------------------------ marks --

    def _resolve_logo(self, logo):
        if logo:
            return logo
        if self.default_logo and Path(self.default_logo).is_file():
            return self.default_logo
        return None

    def org_mark(self, shapes, l, t, h, *, on_dark: bool, logo=None):
        """Left-aligned brand lockup: logo image (or placeholder circle)
        with the org name beside it."""
        logo = self._resolve_logo(logo)
        if not logo and not self.org_lines:
            return
        color = self.on_primary if on_dark else self.primary
        h_in = h / 914400
        if logo:
            src = _logo_white(logo) if on_dark else str(logo)
            pic = shapes.add_picture(src, int(l), int(t), height=int(h))
            if self.logo_has_wordmark:
                return
            x = l + pic.width + IN(0.12)
        else:
            circle(shapes, l + h / 2, t + h / 2, h, fill=color)
            x = l + int(h * 1.25)
        tf = textbox(shapes, Rect(x, t, IN(h_in * 5), h), anchor="m", wrap=False)
        size = max(6.5, h_in * 21)
        for line in self.org_lines:
            run(para(tf, line_spacing=0.98), line, font=self.body_font,
                size=size, color=color, spacing=1.2)

    def org_mark_centered(self, shapes, t, h, *, on_dark: bool, logo=None):
        """Horizontally centered, stacked brand mark for title slides:
        symbol on top, org name below."""
        logo = self._resolve_logo(logo)
        if not logo and not self.org_lines:
            return
        color = self.on_primary if on_dark else self.primary
        h_in = h / 914400
        d = int(h * 0.55)
        if logo:
            src = _logo_white(logo) if on_dark else str(logo)
            if self.logo_has_wordmark:
                pic = shapes.add_picture(src, 0, int(t + h * 0.14), height=int(h * 0.72))
                pic.left = (W - pic.width) // 2
                return
            pic = shapes.add_picture(src, 0, int(t), height=d)
            pic.left = (W - pic.width) // 2
        else:
            circle(shapes, W // 2, t + d // 2, d, fill=color)
        tf = textbox(shapes, Rect(0, t + int(h * 0.62), W, int(h * 0.38)))
        size = max(7.5, h_in * 9.5)
        for line in self.org_lines:
            run(para(tf, align="c", line_spacing=1.0), line, font=self.body_font,
                size=size, color=color, spacing=1.2)

    # ------------------------------------------------------- title slide --

    def _title_canvas(self, shapes):
        """Title-slide background; themes overlay brand motifs here."""
        box(shapes, Rect(0, 0, W, H), fill=self.primary)

    @staticmethod
    def _est_lines(text, size, w_emu):
        """Rough wrap estimate (average glyph ~0.54 em) for flow layout."""
        per_line = max(1, int((w_emu / IN(1)) / (size * 0.0075)))
        return max(1, -(-len(text) // per_line))

    def title_slide(self, slide, *, title, subtitle=None, authors=None, logo=None):
        """Left-anchored composition: the text block flows down from the
        upper third, the brand lockup anchors the bottom-left corner."""
        sh = slide.shapes
        self._title_canvas(sh)

        x, wmax = IN(1.15), W - IN(2.3)
        y = IN(1.85)
        lines = self._est_lines(title, self.t_title, wmax)
        h = int(lines * IN(self.t_title * 1.28 / 72))
        tf = textbox(sh, Rect(x, y, wmax, h))
        run(para(tf, line_spacing=1.04), title, font=self.title_font,
            size=self.t_title, bold=True, color=self.on_primary)
        y += h + IN(0.32)

        if subtitle:
            sub_color = self.title_subtitle_color or \
                _mix(self.on_primary, self.primary, 0.22)
            sub_h = int(self._est_lines(subtitle, 18, wmax) * IN(0.34))
            tf = textbox(sh, Rect(x, y, wmax, sub_h))
            run(para(tf, line_spacing=1.1), subtitle, font=self.body_font,
                size=18, color=sub_color)
            y += sub_h + IN(0.55)

        if authors:
            authors = ", ".join(authors) if not isinstance(authors, str) else authors
            tf = textbox(sh, Rect(x, y, wmax, IN(0.65)))
            run(para(tf, line_spacing=1.15), authors, font=self.body_font,
                size=15, color=self.title_authors_color)

        self.org_mark(sh, x, H - IN(1.18), IN(0.6), on_dark=True, logo=logo)
        return slide

    # ---------------------------------------------------- content chrome --

    def content_chrome(self, slide, *, title, references=None,
                       logo=None) -> Rect:
        """Background, title, footer, logo, refs.
        Returns the rectangle available for slide content."""
        sh = slide.shapes
        box(sh, Rect(0, 0, W, H), fill=self.bg)
        fh = self._footer_height(references)
        ft = H - fh
        box(sh, Rect(0, ft, W, fh), fill=self.primary)

        logo = self._resolve_logo(logo)
        mark_h = self.footer_h - (IN(0.12) if logo else IN(0.24))
        self.org_mark(sh, IN(0.25), ft + (fh - mark_h) // 2, mark_h,
                      on_dark=True, logo=logo)
        if references:
            self._references(sh, ft, fh, references)

        tf = textbox(sh, Rect(IN(self.c_title_x), IN(self.c_title_y),
                              W - IN(self.c_title_x) - self.margin, IN(0.78)))
        run(para(tf), title, font=self.title_font, size=self.c_title,
            bold=True, color=self.primary)

        top = IN(self.c_title_y + 0.83)
        return Rect(self.margin, top, W - 2 * self.margin, ft - IN(0.18) - top)

    def _footer_height(self, references):
        """Footer band; grows only when references need more than two rows."""
        if not references or isinstance(references, str):
            return self.footer_h
        n = len(references)
        if n <= 1:
            return self.footer_h
        cols = 1 if n <= 2 else 2
        rows = -(-n // cols)
        return max(self.footer_h, IN(0.20) + rows * IN(0.205))

    def _references(self, sh, ft, fh, refs):
        """References one per cell (not run together on a line): a single
        stacked column for <=2, two columns for more. The zone is anchored
        to the right edge and sized to the text, so a short list gets a
        small box instead of a slide-wide one. A lone citation is special:
        no [1], a snugger pill, slightly larger text."""
        items = [refs] if isinstance(refs, str) else list(refs)
        single = len(items) == 1
        size = self.refs_single if single else self.refs
        labelled = items if single else \
            [f"[{i + 1}] {r}" for i, r in enumerate(items)]
        cols = 1 if len(items) <= 2 else 2
        gap = IN(0.3)
        col_ws = [max(self._est_text_w(t, size) for t in labelled[c::cols])
                  for c in range(cols)]
        content_w = sum(col_ws) + (cols - 1) * gap
        if self.refs_boxed:
            pad = IN(0.14)
            right = W - IN(0.14)
            left = max(IN(3.05), right - content_w - 2 * pad)
            band_h = IN(0.34) if single else fh - IN(0.18)
            zone = Rect(left, ft + (fh - band_h) // 2, right - left, band_h)
            box(sh, zone, fill=WHITE, line=BLACK, line_w=0.75, rounded=True,
                radius_in=0.05)
            inner = Rect(zone.l + pad, zone.t + IN(0.03),
                         zone.w - 2 * pad, zone.h - IN(0.06))
            self._refs_grid(sh, inner, labelled, cols, col_ws, gap, size, BLACK)
        else:
            right = W - IN(0.25)
            left = max(IN(3.2), right - content_w)
            zone = Rect(left, ft + IN(0.05), right - left, fh - IN(0.10))
            self._refs_grid(sh, zone, labelled, cols, col_ws, gap, size,
                            self.on_primary)

    def _est_text_w(self, text, size=None):
        """Crude width estimate (average glyph ~0.56 em)."""
        return IN(len(str(text)) * (size or self.refs) * 0.0078)

    def _refs_grid(self, sh, zone, labelled, cols, col_ws, gap, size, color):
        """Cells fill left-to-right then down; columns sized to their text
        (scaled down together if the zone had to be clamped)."""
        rows = -(-len(labelled) // cols)
        scale = min(1.0, zone.w / max(1, sum(col_ws) + (cols - 1) * gap))
        xs, x = [], 0
        for wcol in col_ws:
            xs.append(int(x * scale))
            x += wcol + gap
        rh = zone.h // rows
        for idx, text in enumerate(labelled):
            c, r = idx % cols, idx // cols
            cell = Rect(zone.l + xs[c], zone.t + r * rh,
                        int(col_ws[c] * scale), rh)
            tf = textbox(sh, cell, anchor="m", wrap=False)
            run(para(tf), text, font=self.body_font, size=size, color=color)

    # ----------------------------------------------------- outline slide --

    def outline_slide(self, slide, sections, *, current=None, title="Outline",
                      logo=None):
        """Numbered section list; with current=i, every other row is dimmed.

        Dimming is a blend toward the background color (pptx text has no
        alpha), so the reprise reads as the same slide with one row lit.
        """
        area = self.content_chrome(slide, title=title, logo=logo)
        sh = slide.shapes
        n = max(1, len(sections))
        row_h = min(IN(0.95), area.h // n)
        d = min(IN(0.52), int(row_h * 0.62))
        y0 = area.t + (area.h - row_h * n) // 2
        x0 = area.l + IN(0.85)
        for i, name in enumerate(sections):
            dim = current is not None and i != current
            cy = y0 + i * row_h + row_h // 2
            if current == i:
                aw, ah = IN(0.34), IN(0.16)
                arrow(sh, x0 - aw - IN(0.18), cy - ah // 2, aw, ah,
                      fill=self.takeaway_color)
            circle(sh, x0 + d // 2, cy, d,
                   fill=_mix(self.primary, self.bg, 0.85) if dim else self.primary,
                   text=str(i + 1), font=self.body_font, size=15,
                   text_color=_mix(self.primary, self.bg, 0.45) if dim
                   else self.on_primary)
            tf = textbox(sh, Rect(x0 + d + IN(0.3), cy - IN(0.3),
                                  area.r - x0 - d - IN(0.4), IN(0.6)), anchor="m")
            run(para(tf), str(name), font=self.body_font, size=20,
                bold=(current == i),
                color=_mix(self.text, self.bg, 0.72) if dim else self.text)
        return slide

    # ------------------------------------------------------- small parts --

    def caption_par(self, tf, text):
        run(para(tf, align="c"), text, font=self.body_font, size=self.caption, color=self.muted)

    def takeaway_line(self, shapes, r: Rect, text):
        """Bold '→ message' line: solid arrow shape + text."""
        aw, ah = IN(0.36), IN(0.17)
        arrow(shapes, r.l, r.cy - ah // 2, aw, ah, fill=self.takeaway_color)
        tf = textbox(shapes, Rect(r.l + aw + IN(0.14), r.t,
                                  r.w - aw - IN(0.14), r.h), anchor="m")
        run(para(tf), text, font=self.body_font, size=self.takeaway, bold=True,
            color=self.takeaway_color)


@register_theme
class Plain(Theme):
    """Neutral default: deep slate-blue, no logo, no organisation name.
    Pass Deck(org=...) for a text logotype, Deck(logo=...) for an image."""

    name = "plain"
    primary = RGBColor(0x1F, 0x3A, 0x5F)
    accent = RGBColor(0x1F, 0x3A, 0x5F)
    takeaway_color = RGBColor(0x1F, 0x3A, 0x5F)
    title_authors_color = RGBColor(0xDD, 0xE6, 0xF0)
    org_lines = ()

    def _title_canvas(self, shapes):
        """Slate field with a lighter band along the bottom edge."""
        box(shapes, Rect(0, 0, W, H), fill=self.primary)
        box(shapes, Rect(0, H - IN(0.35), W, IN(0.35)),
            fill=mix(self.primary, WHITE, 0.25))
