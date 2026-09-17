"""The Deck: a reusable slide factory on top of python-pptx.

    from slide_factory import Deck

    deck = Deck()                                # neutral "plain" theme
    deck = Deck(theme="my_lab")                  # a theme from an installed theme package
    deck.title_slide(title="...", authors=[...])
    deck.picture_slide("Title", "fig.png", caption="...")
    deck.two_picture_slide("Title", [("a.png", "left"), ("b.png", "right")])
    deck.picture_bullets_slide("Title", "fig.png", bullets=["...", ("sub", 1)])
    deck.save("out.pptx")

Every slide function returns the pptx slide object, so you can post-tweak.
"""
from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation

from .animate import appear_on_click
from .draw import (IN, SLIDE_H, SLIDE_W, Rect, box, bullet, circle,
                   circle_crop, fit_picture, flow_arrow, line, para, run, textbox)
from .themes import Theme, _mix, make_theme

_CAPTION_H = IN(0.34)
_TAKEAWAY_H = IN(0.45)


def _check(picture):
    """Validate a picture argument: file path (str/Path) or open file/BytesIO."""
    if hasattr(picture, "read"):
        return picture
    path = str(picture)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Picture not found: {path}")
    return path


# ---------------------------------------------------- conceptual plumbing --

def _node(sh, th, r: Rect, text, *, style="plain", size=17):
    """Rounded chain node. style: plain | hot (brand primary) | dim."""
    accent = th.accent_color()
    if style == "hot":
        shp = box(sh, r, fill=th.primary, rounded=True, radius_in=0.12)
        color, bold = th.on_primary, True
    elif style == "dim":
        shp = box(sh, r, fill=_mix(th.bg, th.text, 0.05),
                  line=_mix(th.text, th.bg, 0.6), line_w=1.0,
                  rounded=True, radius_in=0.12)
        color, bold = th.muted, False
    else:
        shp = box(sh, r, fill=th.bg, line=accent, line_w=1.5,
                  rounded=True, radius_in=0.12)
        color, bold = accent, True
    tf = shp.text_frame
    tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, IN(0.05))
    run(para(tf, align="c"), text, font=th.body_font, size=size, bold=bold,
        color=color)
    return shp


def _headline(sh, th, area: Rect, big, sub=None):
    """Big centered statement at the top of the content area.
    Returns the y where following content may start."""
    y = area.t + IN(0.1)
    tf = textbox(sh, Rect(area.l, y, area.w, IN(0.75)), anchor="m")
    run(para(tf, align="c"), big, font=th.title_font, size=30, bold=True,
        color=th.accent_color())
    y += IN(0.85)
    if sub:
        tf = textbox(sh, Rect(area.l, y, area.w, IN(0.4)), anchor="m")
        run(para(tf, align="c"), sub, font=th.body_font, size=17, color=th.muted)
        y += IN(0.55)
    return y


def _hchain(sh, th, band: Rect, nodes, *, hot=(), dim=False, node_h=1.05,
            size=17, arrows=True):
    """Horizontal chain of nodes, with or without arrows, centered in band."""
    n = len(nodes)
    gap = IN(0.6 if arrows else 0.35)
    bw = (band.w - gap * (n - 1)) // n
    bh = IN(node_h)
    y = band.t + (band.h - bh) // 2
    for i, text in enumerate(nodes):
        x = band.l + i * (bw + gap)
        style = "hot" if i in hot else ("dim" if dim else "plain")
        _node(sh, th, Rect(x, y, bw, bh), text, style=style, size=size)
        if arrows and i < n - 1:
            flow_arrow(sh, x + bw + gap // 2, y + bh // 2, w=IN(0.42), h=IN(0.2),
                       fill=_mix(th.text, th.bg, 0.55) if dim else th.primary)


def card_item(text, *, sub=None, tag=None, tag_color=None, tag_italic=False,
              muted=False):
    """One entry of a columns_slide card.

    sub: muted sub-point under the entry; tag: short uppercase-style label
    after the text (tag_color defaults to the theme's muted grey, tag_italic
    gives a lighter, non-bold tag); muted: grey italic entry text.
    """
    return {"text": text, "sub": sub, "tag": tag, "tag_color": tag_color,
            "tag_italic": tag_italic, "muted": muted}


class Deck:
    """A presentation with a fixed theme. Feed it content, get slides."""

    def __init__(self, theme: str | Theme = "plain", logo=None, org=None):
        self.theme = make_theme(theme)
        self.logo = logo  # optional path to a real logo image, used everywhere
        if org:           # text logotype: a name, or a sequence of stacked lines
            self.theme.org_lines = (org,) if isinstance(org, str) else tuple(org)
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self._outline_sections = None   # set by outline(); used by next_section()
        self._outline_title = "Outline"
        self._section_i = 0

    # ------------------------------------------------------------ slides --

    def title_slide(self, title: str, *, subtitle: str | None = None,
                    authors=None, logo=None, notes: str | None = None):
        """Big branded opening slide."""
        s = self._blank()
        self.theme.title_slide(s, title=title, subtitle=subtitle, authors=authors,
                               logo=logo or self.logo)
        return self._done(s, notes)

    def outline(self, sections, *, title: str = "Outline",
                notes: str | None = None):
        """Talk outline: numbered section list, all rows bright.

        Remembers the sections; call next_section() before each section of
        the talk to reprise this slide with that section highlighted and the
        rest dimmed.
        """
        self._outline_sections = [str(s) for s in sections]
        self._outline_title = title
        self._section_i = 0
        s = self._blank()
        self.theme.outline_slide(s, self._outline_sections, current=None,
                                 title=title, logo=self.logo)
        return self._done(s, notes)

    def next_section(self, *, notes: str | None = None):
        """Outline reprise: the next section lit, everything else dimmed."""
        if not self._outline_sections:
            raise RuntimeError("call outline([...]) before next_section()")
        if self._section_i >= len(self._outline_sections):
            raise RuntimeError(
                f"next_section() called more than {len(self._outline_sections)} "
                "times (one per outline section)")
        s = self._blank()
        self.theme.outline_slide(s, self._outline_sections,
                                 current=self._section_i,
                                 title=self._outline_title, logo=self.logo)
        self._section_i += 1
        return self._done(s, notes)

    def picture_slide(self, title: str, picture, *, caption: str | None = None,
                      takeaway: str | None = None, references=None,
                      notes: str | None = None):
        """One dominant, centered picture.

        picture: file path (str/Path) or open binary file / BytesIO.
        Formats: PNG, JPEG, GIF, BMP, TIFF (animated GIFs stay animated
        in slideshow mode). Same for all picture args on other slides.
        """
        s = self._blank()
        area = self._chrome(s, title, references)
        area = self._takeaway(s, area, takeaway)
        self._captioned_picture(s, _check(picture), area, caption)
        return self._done(s, notes)

    def two_picture_slide(self, title: str, pictures, *, takeaway: str | None = None,
                          references=None, notes: str | None = None):
        """Two pictures side by side.

        pictures: sequence of exactly 2 items, each a path or (path, caption).
        """
        if len(pictures) != 2:
            raise ValueError("two_picture_slide needs exactly 2 pictures")
        s = self._blank()
        area = self._chrome(s, title, references)
        area = self._takeaway(s, area, takeaway)
        gap = IN(0.5)
        cell_w = (area.w - gap) // 2
        for i, item in enumerate(pictures):
            path, cap = item if isinstance(item, (tuple, list)) else (item, None)
            cell = Rect(area.l + i * (cell_w + gap), area.t, cell_w, area.h)
            self._captioned_picture(s, _check(path), cell, cap)
        return self._done(s, notes)

    def picture_bullets_slide(self, title: str, picture, bullets, *,
                              caption: str | None = None,
                              bullets_heading: str | None = None,
                              takeaway: str | None = None, references=None,
                              picture_side: str = "right",
                              notes: str | None = None):
        """Bullet points next to a picture.

        bullets: sequence of str or (str, level) — level 0 is top level.
        picture_side: "right" (default) or "left".
        """
        s = self._blank()
        area = self._chrome(s, title, references)
        area = self._takeaway(s, area, takeaway)

        gap = IN(0.45)
        text_w = int(area.w * 0.44)
        pic_w = area.w - text_w - gap
        if picture_side == "right":
            t_rect = Rect(area.l, area.t, text_w, area.h)
            p_rect = Rect(area.l + text_w + gap, area.t, pic_w, area.h)
        elif picture_side == "left":
            p_rect = Rect(area.l, area.t, pic_w, area.h)
            t_rect = Rect(area.l + pic_w + gap, area.t, text_w, area.h)
        else:
            raise ValueError("picture_side must be 'left' or 'right'")

        th = self.theme
        y = t_rect.t
        if bullets_heading:
            tf = textbox(s.shapes, Rect(t_rect.l, y, t_rect.w, IN(0.42)))
            run(para(tf), bullets_heading, font=th.title_font, size=18,
                bold=True, color=th.text)
            y += IN(0.5)
        tf = textbox(s.shapes, Rect(t_rect.l, y, t_rect.w, t_rect.b - y))
        for item in bullets:
            text, level = item if isinstance(item, (tuple, list)) else (item, 0)
            p = para(tf, space_after=7, line_spacing=1.12)
            bullet(p, char=th.bullet_char if level == 0 else th.sub_bullet_char,
                   font=th.body_font, level=level)
            run(p, text, font=th.body_font,
                size=th.body if level == 0 else th.sub,
                color=th.text if level == 0 else th.muted)

        self._captioned_picture(s, _check(picture), p_rect, caption)
        return self._done(s, notes)

    def picture_stack_slide(self, title: str, pictures, *, takeaway: str | None = None,
                            references=None, reveal_first: bool = False,
                            notes: str | None = None):
        """Pictures stacked in the same spot, revealed one per click (build-up).

        In a slideshow, each click makes the next picture (and its caption)
        appear on top of the previous one; the editing view shows the full
        stack. Works best when the pictures share one size/aspect.

        pictures: sequence of items, each a path or (path, caption).
        reveal_first: if True, even the first picture waits for a click.
        """
        s = self._blank()
        area = self._chrome(s, title, references)
        area = self._takeaway(s, area, takeaway)
        any_caption = any(isinstance(i, (tuple, list)) and i[1] for i in pictures)
        cap_h = _CAPTION_H if any_caption else 0
        pic_area = Rect(area.l, area.t, area.w, area.h - cap_h)

        groups = []
        for k, item in enumerate(pictures):
            path, cap = item if isinstance(item, (tuple, list)) else (item, None)
            group = []
            if k > 0:  # backing sheet so each stage cleanly covers the last
                group.append(box(s.shapes, area, fill=self.theme.bg))
            group.append(fit_picture(s.shapes, _check(path), pic_area))
            if cap:
                tf = textbox(s.shapes, Rect(area.l, area.b - IN(0.3), area.w, IN(0.3)))
                self.theme.caption_par(tf, cap)
                group.append(tf._parent)
            groups.append(group)

        appear_on_click(s, groups if reveal_first else groups[1:])
        return self._done(s, notes)

    # ------------------------------------------------------- text slides --

    def bullets_slide(self, title, bullets, *, headline=None, sub=None,
                      references=None, notes=None):
        """Full-width bullet list under an optional big headline.

        bullets: str or (lead, rest) — a bold accent-colored lead-in word,
        then the text. Without a headline the list is vertically centered.
        """
        s = self._blank()
        th = self.theme
        area = self._chrome(s, title, references)
        y = area.t + IN(0.1)
        if headline:
            y = _headline(s.shapes, th, area, headline, sub) + IN(0.15)
        tf = textbox(s.shapes, Rect(area.l + IN(0.9), y, area.w - IN(1.8),
                                    area.b - y - IN(0.3)),
                     anchor="t" if headline else "m")
        for item in bullets:
            lead, text = item if isinstance(item, (tuple, list)) else (None, item)
            p = para(tf, space_after=17, line_spacing=1.12)
            bullet(p, char=th.bullet_char, font=th.body_font)
            if lead:
                run(p, lead + "  ", font=th.body_font, size=22, bold=True,
                    color=th.accent_color())
            run(p, text, font=th.body_font, size=22, color=th.text)
        return self._done(s, notes)

    def columns_slide(self, title, columns, *, size=16, card_h=None, weights=None,
                      legend=None, references=None, notes=None):
        """Side-by-side cards of bulleted entries.

        columns: (heading, subheading_or_None, items) per card; an item is a
        str or a card_item(...) with sub-point, colored tag and muted styling.
        size: entry font size; card_h: card height in inches (vertically
        centered; default the full content area); weights: relative card
        widths; legend: a muted line under the cards, or a list of
        (label, color) for a colored key.
        """
        s = self._blank()
        th = self.theme
        accent = th.accent_color()
        area = self._chrome(s, title, references)
        foot = IN(0.45) if legend else 0
        n = len(columns)
        gap = IN(0.35)
        weights = weights or [1] * n
        free = area.w - gap * (n - 1)
        widths = [int(free * w / sum(weights)) for w in weights]
        full = area.h - IN(0.2) - foot
        ch = min(IN(card_h), full) if card_h else full
        for i, (heading, subheading, items) in enumerate(columns):
            left = area.l + sum(widths[:i]) + i * gap
            r = Rect(left, area.t + IN(0.1) + (full - ch) // 2, widths[i], ch)
            box(s.shapes, r, fill=_mix(th.bg, accent, 0.04), line=accent,
                line_w=1.5, rounded=True, radius_in=0.18)
            tf = textbox(s.shapes, r.inset(0.3, 0.28))
            run(para(tf, space_after=2 if subheading else 12), heading,
                font=th.title_font, size=19, bold=True, color=accent)
            if subheading:
                run(para(tf, space_after=12), subheading, font=th.body_font,
                    size=13, italic=True, color=th.muted)
            for item in items:
                if isinstance(item, str):
                    item = card_item(item)
                p = para(tf, space_after=3 if item["sub"] else 9, line_spacing=1.1)
                bullet(p, char=th.bullet_char, font=th.body_font)
                run(p, item["text"], font=th.body_font, size=size,
                    italic=item["muted"],
                    color=th.muted if item["muted"] else th.text)
                if item["tag"]:  # non-breaking, so a tag never splits over lines
                    run(p, "   " + item["tag"].replace(" ", "\u00a0"),
                        font=th.body_font, size=size - 4,
                        bold=not item["tag_italic"], italic=item["tag_italic"],
                        color=item["tag_color"] or th.muted)
                if item["sub"]:
                    p = para(tf, space_after=9, line_spacing=1.08)
                    bullet(p, char=th.sub_bullet_char, font=th.body_font, level=1)
                    run(p, item["sub"], font=th.body_font, size=size - 3,
                        color=th.muted)
        if legend:
            tf = textbox(s.shapes, Rect(area.l, area.b - IN(0.38), area.w, IN(0.35)),
                         anchor="m")
            p = para(tf, align="c")
            if isinstance(legend, str):
                run(p, legend, font=th.body_font, size=14, italic=True,
                    color=th.muted)
            else:
                for k, (label, color) in enumerate(legend):
                    if k:
                        run(p, "      ", font=th.body_font, size=13)
                    run(p, "\u25cf " + label, font=th.body_font, size=13,
                        bold=True, color=color)
        return self._done(s, notes)

    def timeline_slide(self, title, months, milestones, spans, *, subtitle=None,
                       aside=None, footnote=None, notes=None):
        """Month axis with deadline milestones above and work spans below.

        months: tick labels, one per month from the axis start.
        milestones: (position_in_months, date_label, text) — boxes above.
        spans: (start, end, text, hot) in months — one bar per row below.
        subtitle: muted line under the title; aside: (start, end, text) — a
        dimmed note box below the axis, for what is outside the plan.
        """
        s = self._blank()
        th = self.theme
        sh = s.shapes
        area = self._chrome(s, title, None)
        if subtitle:
            tf = textbox(sh, Rect(area.l, area.t - IN(0.12), area.w, IN(0.4)))
            run(para(tf), subtitle, font=th.body_font, size=18, color=th.muted)
            area = Rect(area.l, area.t + IN(0.35), area.w, area.h - IN(0.35))
        foot = IN(0.5) if footnote else 0
        ax = Rect(area.l + IN(0.3), area.t, area.w - IN(0.6), area.h - foot)
        bw, bh, stem = IN(2.75), IN(1.0), IN(0.8)
        row_h, row_gap = IN(0.6), IN(0.18)
        below = IN(0.3) + len(spans) * (row_h + row_gap)
        above = stem + bh + IN(0.4)
        y = ax.t + above + (ax.h - above - below) // 2
        unit = ax.w / len(months)
        grey = _mix(th.text, th.bg, 0.6)

        def x_of(m):
            return int(ax.l + m * unit)

        line(sh, ax.l, y, ax.r, y, color=grey, w=2.5)
        for i, label in enumerate(months):
            line(sh, x_of(i), y - IN(0.08), x_of(i), y + IN(0.08), color=grey, w=1.5)
            tf = textbox(sh, Rect(x_of(i), y - IN(0.4), int(unit), IN(0.3)))
            run(para(tf, align="c"), label, font=th.body_font, size=13,
                color=th.muted)
        line(sh, ax.r, y - IN(0.08), ax.r, y + IN(0.08), color=grey, w=1.5)

        for k, (start, end, text, hot) in enumerate(spans):
            top = y + IN(0.3) + k * (row_h + row_gap)
            _node(sh, th, Rect(x_of(start), top, x_of(end) - x_of(start), row_h),
                  text, style="hot" if hot else "plain", size=17)

        for m, when, what in milestones:
            x = x_of(m)
            top = y - stem - bh
            line(sh, x, y, x, top + bh, color=th.primary, w=2.0)
            circle(sh, x, y, IN(0.22), fill=th.primary)
            left = min(max(x - bw // 2, area.l), area.r - bw)
            _node(sh, th, Rect(left, top, bw, bh), what, size=15)
            tf = textbox(sh, Rect(left, top - IN(0.38), bw, IN(0.32)), anchor="m")
            run(para(tf, align="c"), when, font=th.body_font, size=16, bold=True,
                color=th.primary)

        if aside:
            start, end, text = aside
            _node(sh, th, Rect(x_of(start), y + IN(0.3), x_of(end) - x_of(start),
                               2 * row_h + row_gap), text, style="dim", size=16)

        if footnote:
            tf = textbox(sh, Rect(area.l, area.b - IN(0.4), area.w, IN(0.38)),
                         anchor="m")
            run(para(tf, align="c"), footnote, font=th.body_font, size=14,
                italic=True, color=th.muted)
        return self._done(s, notes)

    # ------------------------------------------------- conceptual slides --

    def chain_slide(self, title, nodes, *, hot=(), headline=None, sub=None,
                    footnote=None, arrows=True, notes=None):
        """Headline + one horizontal row of nodes (+ optional footnote line).

        nodes: strings; hot: indices drawn in the brand primary;
        arrows=False gives a plain row (e.g. a list of cases).
        """
        s = self._blank()
        th = self.theme
        area = self._chrome(s, title, None)
        top = _headline(s.shapes, th, area, headline, sub) if headline else area.t
        foot_h = IN(0.55) if footnote else 0
        band = Rect(area.l + IN(0.3), top, area.w - IN(0.6), area.b - top - foot_h)
        _hchain(s.shapes, th, band, nodes, hot=hot,
                size=18 if len(nodes) <= 3 else 16, arrows=arrows)
        if footnote:
            tf = textbox(s.shapes, Rect(area.l, area.b - IN(0.45), area.w, IN(0.4)))
            run(para(tf, align="c"), footnote, font=th.body_font, size=15,
                color=th.muted, italic=True)
        return self._done(s, notes)

    def vchain_slide(self, title, nodes, *, hot=(), headline=None, notes=None):
        """Headline + vertical flow chain."""
        s = self._blank()
        th = self.theme
        area = self._chrome(s, title, None)
        top = _headline(s.shapes, th, area, headline) if headline else area.t
        n = len(nodes)
        gap = IN(0.42)
        bh = min(IN(0.8), (area.b - top - gap * (n - 1)) // n)
        bw = IN(4.6)
        x = area.l + (area.w - bw) // 2
        total = n * bh + (n - 1) * gap
        y = top + (area.b - top - total) // 2
        for i, text in enumerate(nodes):
            _node(s.shapes, th, Rect(x, y, bw, bh), text,
                  style="hot" if i in hot else "plain", size=17)
            if i < n - 1:
                flow_arrow(s.shapes, x + bw // 2, y + bh + gap // 2,
                           w=IN(0.2), h=IN(0.34), fill=th.primary,
                           direction="down")
            y += bh + gap
        return self._done(s, notes)

    def labeled_chains_slide(self, title, rows, *, headline=None, notes=None):
        """Rows of (label, nodes, kwargs) chains, e.g. old way vs new way.
        kwargs per row: hot=(indices,), dim=True."""
        s = self._blank()
        th = self.theme
        area = self._chrome(s, title, None)
        top = _headline(s.shapes, th, area, headline) if headline else area.t
        row_h = (area.b - top) // len(rows)
        for i, (label, nodes, kw) in enumerate(rows):
            y = top + i * row_h
            tf = textbox(s.shapes, Rect(area.l + IN(0.1), y, IN(2.0), row_h),
                         anchor="m")
            run(para(tf), label, font=th.body_font, size=15, bold=True,
                color=th.muted if kw.get("dim") else th.accent_color())
            band = Rect(area.l + IN(2.3), y + IN(0.12), area.w - IN(2.5),
                        row_h - IN(0.24))
            _hchain(s.shapes, th, band, nodes, hot=kw.get("hot", ()),
                    dim=kw.get("dim", False),
                    node_h=min(0.9, (row_h - IN(0.24)) / IN(1)),
                    size=15 if len(nodes) > 3 else 16)
        return self._done(s, notes)

    def triangle_slide(self, title, corners, *, bottom=None, bottom_sub=None,
                       notes=None):
        """Three nodes joined into a triangle, plus a bottom statement."""
        from .draw import line as _draw_line

        s = self._blank()
        th = self.theme
        area = self._chrome(s, title, None)
        foot = IN(1.45) if bottom else 0
        zone = Rect(area.l + IN(2.2), area.t + IN(0.05), area.w - IN(4.4),
                    area.h - foot)
        bw, bh = IN(2.9), IN(0.85)
        top_c = (zone.cx, zone.t + bh // 2)
        bl_c = (zone.l + bw // 2, zone.b - bh // 2)
        br_c = (zone.r - bw // 2, zone.b - bh // 2)
        for a, b in ((top_c, bl_c), (top_c, br_c), (bl_c, br_c)):
            _draw_line(s.shapes, a[0], a[1], b[0], b[1], color=th.primary, w=2.2)
        for (cx, cy), text in zip((top_c, bl_c, br_c), corners):
            _node(s.shapes, th, Rect(cx - bw // 2, cy - bh // 2, bw, bh), text,
                  style="plain", size=17)
        if bottom:
            tf = textbox(s.shapes, Rect(area.l, area.b - IN(0.95), area.w, IN(0.9)))
            run(para(tf, align="c", space_after=4), bottom, font=th.title_font,
                size=20, bold=True, color=th.primary)
            if bottom_sub:
                run(para(tf, align="c"), bottom_sub, font=th.body_font, size=14,
                    color=th.muted)
        return self._done(s, notes)

    def two_cards_slide(self, title, cards, *, notes=None):
        """Two big numbered cards side by side: (heading, [lines])."""
        s = self._blank()
        th = self.theme
        accent = th.accent_color()
        area = self._chrome(s, title, None)
        gap = IN(0.6)
        cw = (area.w - gap) // 2
        ch = area.h - IN(0.6)
        for i, (heading, lines) in enumerate(cards):
            r = Rect(area.l + i * (cw + gap), area.t + IN(0.2), cw, ch)
            box(s.shapes, r, fill=_mix(th.bg, accent, 0.04), line=accent,
                line_w=1.5, rounded=True, radius_in=0.18)
            circle(s.shapes, r.l + IN(0.75), r.t + IN(0.85), IN(0.75),
                   fill=th.primary, text=str(i + 1), font=th.title_font, size=26)
            tf = textbox(s.shapes, Rect(r.l + IN(0.55), r.t + IN(1.6),
                                        r.w - IN(1.1), r.h - IN(2.0)))
            run(para(tf, space_after=14), heading, font=th.title_font, size=23,
                bold=True, color=accent)
            for text in lines:
                run(para(tf, space_after=8, line_spacing=1.15), text,
                    font=th.body_font, size=16, color=th.text)
        return self._done(s, notes)

    def team_slide(self, title, people, *, closing=None, notes=None):
        """Row of circular headshots with names; entries without a photo get
        an initials medallion. people: (photo_or_None, name, affiliation)."""
        s = self._blank()
        th = self.theme
        area = self._chrome(s, title, None)
        n = len(people)
        cell_w = area.w // n
        d = min(IN(1.7), int(cell_w * 0.8))
        top = area.t + IN(0.55)
        for i, (photo, name, affil) in enumerate(people):
            cx = area.l + i * cell_w + cell_w // 2
            if photo:
                s.shapes.add_picture(circle_crop(_check(photo)),
                                     int(cx - d / 2), int(top), int(d), int(d))
            else:
                initials = "".join(w[0] for w in name.split()[:2])
                circle(s.shapes, cx, top + d // 2, d, fill=th.accent_color(),
                       text=initials, font=th.title_font, size=30)
            tf = textbox(s.shapes, Rect(area.l + i * cell_w, top + d + IN(0.18),
                                        cell_w, IN(0.95)))
            run(para(tf, align="c"), name, font=th.body_font, size=14, bold=True,
                color=th.text)
            run(para(tf, align="c"), affil, font=th.body_font, size=11.5,
                color=th.muted)
        if closing:
            tf = textbox(s.shapes, Rect(area.l, area.b - IN(0.85), area.w, IN(0.6)),
                         anchor="m")
            run(para(tf, align="c"), closing, font=th.title_font, size=26,
                bold=True, color=th.primary)
        return self._done(s, notes)

    def group_stack_slide(self, title, steps, *, takeaway=None, references=None,
                          reveal_first=False, label_h_in=0.38, notes=None):
        """Click-through builds where each step reveals a GROUP of pictures
        plus an optional label, cleanly covering the previous step.

        Generalizes picture_stack_slide from one picture per click to a
        laid-out set per click (e.g. map + wave function + histogram that
        must update together).

        steps: sequence of (label_or_None, cells); cells is a sequence of
        (picture, (fx, fy, fw, fh)) with fractional geometry inside the
        picture zone (the content area minus the label strip).
        reveal_first: if True, even the first group waits for a click.
        """
        s = self._blank()
        th = self.theme
        area = self._chrome(s, title, references)
        area = self._takeaway(s, area, takeaway)
        cap_h = IN(label_h_in) if any(label for label, _ in steps) else 0
        zone = Rect(area.l, area.t, area.w, area.h - cap_h)
        groups = []
        for k, (label, cells) in enumerate(steps):
            g = []
            if k:  # backing sheet so each step cleanly covers the last
                g.append(box(s.shapes, area, fill=th.bg))
            for pic, (fx, fy, fw, fh) in cells:
                cell = Rect(zone.l + int(zone.w * fx), zone.t + int(zone.h * fy),
                            int(zone.w * fw), int(zone.h * fh))
                g.append(fit_picture(s.shapes, _check(pic), cell))
            if label:
                tf = textbox(s.shapes, Rect(area.l, area.b - cap_h + IN(0.04),
                                            area.w, cap_h), anchor="m")
                run(para(tf, align="c"), label, font=th.body_font, size=15,
                    bold=True, color=th.primary)
                g.append(tf._parent)
            groups.append(g)
        appear_on_click(s, groups if reveal_first else groups[1:])
        return self._done(s, notes)

    # ------------------------------------------------------------- utils --

    def save(self, path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path

    def _blank(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def _chrome(self, s, title, references) -> Rect:
        return self.theme.content_chrome(s, title=title, references=references,
                                         logo=self.logo)

    def _takeaway(self, s, area: Rect, takeaway) -> Rect:
        """Reserve a bottom strip for the '→ takeaway' line, if any."""
        if not takeaway:
            return area
        self.theme.takeaway_line(s.shapes, Rect(area.l + IN(0.1), area.b - _TAKEAWAY_H,
                                                area.w - IN(0.2), _TAKEAWAY_H), takeaway)
        return Rect(area.l, area.t, area.w, area.h - _TAKEAWAY_H - IN(0.12))

    def _captioned_picture(self, s, path: str, area: Rect, caption):
        cap_h = _CAPTION_H if caption else 0
        pic = fit_picture(s.shapes, path, Rect(area.l, area.t, area.w, area.h - cap_h))
        if caption:
            top = min(pic.top + pic.height + IN(0.06), area.b - IN(0.28))
            tf = textbox(s.shapes, Rect(area.l, top, area.w, IN(0.3)))
            self.theme.caption_par(tf, caption)

    def _done(self, s, notes):
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s
