"""Low-level python-pptx drawing helpers shared by all themes.

Everything here is theme-agnostic plumbing: rectangles, text boxes,
runs, bullets, aspect-ratio-preserving picture placement.
"""
from __future__ import annotations

from PIL import Image as _PILImage
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU_PER_IN = 914400

# Exact 16:9 slide size used by PowerPoint (13.333in x 7.5in)
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)


def IN(inches: float) -> Emu:
    """Inches -> Emu (int)."""
    return Emu(int(round(inches * EMU_PER_IN)))


class Rect:
    """Simple rectangle in EMU with a few conveniences."""

    __slots__ = ("l", "t", "w", "h")

    def __init__(self, l, t, w, h):
        self.l, self.t, self.w, self.h = int(l), int(t), int(w), int(h)

    @property
    def r(self):
        return self.l + self.w

    @property
    def b(self):
        return self.t + self.h

    @property
    def cx(self):
        return self.l + self.w // 2

    @property
    def cy(self):
        return self.t + self.h // 2

    def inset(self, dx: float, dy: float | None = None) -> "Rect":
        """Shrink by dx (and dy) inches on every side."""
        dy = dx if dy is None else dy
        return Rect(self.l + IN(dx), self.t + IN(dy), self.w - 2 * IN(dx), self.h - 2 * IN(dy))

    def __repr__(self):
        f = EMU_PER_IN
        return f"Rect({self.l / f:.2f}, {self.t / f:.2f}, {self.w / f:.2f}, {self.h / f:.2f} in)"


# ---------------------------------------------------------------- shapes ---

def box(shapes, r: Rect, *, fill: RGBColor | None = None, line: RGBColor | None = None,
        line_w: float = 1.0, rounded: bool = False, radius_in: float = 0.1):
    """Rectangle / rounded rectangle. fill=None -> transparent, line=None -> no outline."""
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = shapes.add_shape(kind, r.l, r.t, r.w, r.h)
    if rounded:
        small_in = min(r.w, r.h) / EMU_PER_IN
        try:
            shp.adjustments[0] = max(0.0, min(0.5, radius_in / small_in))
        except (IndexError, ValueError):
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def circle(shapes, cx, cy, d, *, fill: RGBColor | None = None, line: RGBColor | None = None,
           line_w: float = 1.0, text: str | None = None, text_color: RGBColor = WHITE,
           font: str = "Arial", size: float = 18, bold: bool = True):
    """Circle of diameter d centered on (cx, cy), optionally with centered text."""
    shp = shapes.add_shape(MSO_SHAPE.OVAL, int(cx - d / 2), int(cy - d / 2), int(d), int(d))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if text:
        tf = shp.text_frame
        tf.word_wrap = False
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, text, font=font, size=size, color=text_color, bold=bold)
    return shp


def arrow(shapes, l, t, w, h, *, fill: RGBColor):
    """Solid right-arrow autoshape (the '→' of takeaway lines)."""
    shp = shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, int(l), int(t), int(w), int(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = 0.45   # shaft thickness
        shp.adjustments[1] = 0.55   # head length
    except (IndexError, ValueError):
        pass
    return shp


def flow_arrow(shapes, cx, cy, *, w, h, fill: RGBColor, direction: str = "right"):
    """Connector arrow between flow-chain nodes, centered on (cx, cy).
    direction: "right" or "down". Default autoshape proportions (unlike
    arrow(), which is tuned for the takeaway line)."""
    kind = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW}[direction]
    shp = shapes.add_shape(kind, int(cx - w / 2), int(cy - h / 2), int(w), int(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def line(shapes, x1, y1, x2, y2, *, color: RGBColor, w: float = 1.5):
    """Straight line connector."""
    from pptx.enum.shapes import MSO_CONNECTOR

    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    c.shadow.inherit = False
    return c


def set_fill_alpha(shp, frac: float):
    """Make a shape's solid fill translucent (frac = opacity, 0..1)."""
    sf = shp._element.spPr.find(qn("a:solidFill"))
    if sf is None or len(sf) == 0:
        return shp
    clr = sf[0]
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(frac * 100000))}))
    return shp


def set_picture_alpha(pic, frac: float):
    """Make a picture translucent (frac = opacity, 0..1)."""
    blip = pic._element.find(qn("p:blipFill")).find(qn("a:blip"))
    blip.append(blip.makeelement(qn("a:alphaModFix"),
                                 {"amt": str(int(frac * 100000))}))
    return pic


# ------------------------------------------------------------------ text ---

_ANCHORS = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
_ALIGNS = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}


def textbox(shapes, r: Rect, *, anchor: str = "t", wrap: bool = True):
    """Zero-margin text box; returns its text frame."""
    tb = shapes.add_textbox(r.l, r.t, r.w, r.h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = _ANCHORS[anchor]
    return tf


def para(tf, *, align: str = "l", space_after: float = 0.0, space_before: float = 0.0,
         line_spacing: float | None = None):
    """Next paragraph of a text frame (reuses the implicit empty first one)."""
    first = tf.paragraphs[0]
    p = first if (len(tf.paragraphs) == 1 and not first.runs) else tf.add_paragraph()
    p.alignment = _ALIGNS[align]
    if space_after:
        p.space_after = Pt(space_after)
    if space_before:
        p.space_before = Pt(space_before)
    if line_spacing:
        p.line_spacing = line_spacing
    return p


def run(p, text: str, *, font: str = "Arial", size: float = 14, color: RGBColor | None = None,
        bold: bool = False, italic: bool = False, spacing: float | None = None):
    """Add a styled run. spacing = letter spacing in points."""
    r = p.add_run()
    r.text = text
    f = r.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    if color is not None:
        f.color.rgb = color
    if spacing is not None:
        r.font._rPr.set("spc", str(int(spacing * 100)))
    return r


def bullet(p, *, char: str = "•", font: str | None = None, level: int = 0,
           hang_in: float = 0.24):
    """Turn a paragraph into a proper hanging-indent bullet (real buChar, not a
    literal glyph, so wrapped lines align)."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(IN(hang_in) * (level + 1))))
    pPr.set("indent", str(-int(IN(hang_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if font:
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": font}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


# -------------------------------------------------------------- pictures ---

def fit_picture(shapes, source, r: Rect, *, align: str = "c", valign: str = "m"):
    """Place a picture inside r preserving aspect ratio (letterboxed).

    source: file path (str/Path) or open binary file / BytesIO.
    Any format python-pptx supports (PNG, JPEG, GIF, BMP, TIFF, WMF).
    Animated GIFs are embedded byte-for-byte, so they stay animated in
    PowerPoint slideshow mode (viewers show the first frame).
    """
    if hasattr(source, "read"):          # file-like: measure, then rewind
        with _PILImage.open(source) as im:
            iw, ih = im.size
        source.seek(0)
    else:
        source = str(source)
        with _PILImage.open(source) as im:
            iw, ih = im.size
    s = min(r.w / iw, r.h / ih)
    nw, nh = int(iw * s), int(ih * s)
    dx = {"l": 0, "c": (r.w - nw) // 2, "r": r.w - nw}[align]
    dy = {"t": 0, "m": (r.h - nh) // 2, "b": r.h - nh}[valign]
    return shapes.add_picture(source, r.l + dx, r.t + dy, nw, nh)


def circle_crop(source):
    """Square-crop a picture and mask it to a circle; returns a PNG BytesIO
    (transparent corners). For round headshots and similar."""
    import io

    im = _PILImage.open(source).convert("RGBA")
    n = min(im.size)
    im = im.crop(
        ((im.width - n) // 2, (im.height - n) // 2, (im.width + n) // 2, (im.height + n) // 2)
    ).resize((600, 600), _PILImage.LANCZOS)
    from PIL import ImageDraw

    mask = _PILImage.new("L", (600, 600), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, 600, 600), fill=255)
    im.putalpha(mask)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


def formula(lines, *, figsize=(7.6, 4.6), dpi=300):
    """Typeset mathtext to a transparent PNG BytesIO for picture slots.

    lines: sequence of (tex, fontsize, color) tuples, stacked top to bottom
    and centered. Requires matplotlib (imported lazily).
    """
    import io

    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig = plt.figure(figsize=figsize)
    fig.patch.set_alpha(0)
    n = len(lines)
    ys = [0.5] if n == 1 else [0.94 - i * (0.86 / (n - 1)) for i in range(n)]
    for y, (tex, size, color) in zip(ys, lines):
        fig.text(0.5, y, tex, ha="center", va="center", fontsize=size, color=str(color))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, transparent=True, bbox_inches="tight", pad_inches=0.15)
    plt.close(fig)
    buf.seek(0)
    return buf
